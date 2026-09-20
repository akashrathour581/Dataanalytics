import io
import json
import time
import unittest
import zipfile
from unittest.mock import patch
import numpy as np
import pandas as pd
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from server.main import app, rate_windows
from server.analytics_engine import store, generate_automated_eda, sanitize_val, query_aggregation, export_dataframe_bytes
from server.storage import DatasetStore, parse_uploaded_file
from server.cleaning import transform

class EngineTests(unittest.TestCase):
    def profile(self, df): return generate_automated_eda(df, 'test', 'test.csv', 'CSV Data')
    def test_statistics(self):
        p=self.profile(pd.DataFrame({'value':[1,2,3,4,5]}))['column_profiles'][0]
        self.assertEqual((p['mean'],p['median'],p['variance'],p['iqr']),(3,3,2.5,2))
    def test_json_safe_edges(self):
        for df in [pd.DataFrame({'x':[np.nan]}),pd.DataFrame({'x':[1]}),pd.DataFrame({'date':[pd.NaT]}),pd.DataFrame({'x':[]}),pd.DataFrame({'a':[True,False]}),pd.DataFrame({'x':[np.inf,-np.inf]})]:
            json.dumps(self.profile(df),allow_nan=False)
        self.assertIsNone(sanitize_val(pd.NA));self.assertEqual(sanitize_val(1.123456789),1.123456789)
    def test_negative_share_not_invented(self):
        data=self.profile(pd.DataFrame({'group':['a','b'],'value':[-10,1]}))
        self.assertFalse(any('commands' in s for s in data['insights']))
        self.assertFalse(any(c['type']=='pie' for c in data['charts']))
    def test_null_correlation(self):
        data=self.profile(pd.DataFrame({'x':[1,1],'y':[2,3]}))
        self.assertIsNone(data['correlation']['matrix'][0]['value'])
    def test_value_column_query(self):
        self.assertEqual(query_aggregation(pd.DataFrame({'value':[1,1,2]}),'value',agg_func='count')[0],{'key':'1','value':2})
        with self.assertRaises(ValueError):query_aggregation(pd.DataFrame({'x':[1]}),'x','missing')
    def test_clean_preserves_nulls_and_original(self):
        df=pd.DataFrame({'name':[' A ',None]})
        cleaned=transform(df,[{'kind':'trim','column':'name'}])
        self.assertEqual(cleaned.iloc[0,0],'A');self.assertTrue(pd.isna(cleaned.iloc[1,0]));self.assertEqual(df.iloc[0,0],' A ')
    def test_clean_operations(self):
        df=pd.DataFrame({'x':[1.,np.nan,3.],'name':[' a b ',' c d ',None]})
        self.assertEqual(transform(df,[{'kind':'fill','column':'x','method':'mean'}]).iloc[1,0],2)
        self.assertEqual(transform(df,[{'kind':'fill','column':'x','method':'median'}]).iloc[1,0],2)
        self.assertEqual(transform(df,[{'kind':'upper','column':'name'}]).iloc[0,1],' A B ')
        self.assertIn('name_2',transform(df,[{'kind':'split','column':'name','value':' '}]))
        self.assertIn('joined',transform(df,[{'kind':'merge','columns':['x','name'],'replacement':'joined','value':'-'}]))
        self.assertEqual(len(transform(df,[{'kind':'drop_rows','column':'x','value':'1.0'}])),2)
        self.assertEqual(len(transform(pd.DataFrame({'x':[1,1]}),[{'kind':'duplicates'}])),1)
    def test_empty_and_collision(self):
        df=pd.DataFrame({'x':[' ',None],'y':[1,None]})
        self.assertEqual(len(transform(df,[{'kind':'empty_rows'}])),1)
        self.assertNotIn('x',transform(df,[{'kind':'empty_columns'}]))
        with self.assertRaises(ValueError):transform(pd.DataFrame({'A A':[1],'a_a':[2]}),[{'kind':'column_names'}])
    def test_outliers(self):
        df=pd.DataFrame({'x':[1,2,2,3,100]})
        self.assertEqual(transform(df,[{'kind':'detect_outliers','column':'x'}])['x_outlier'].sum(),1)
        self.assertEqual(len(transform(df,[{'kind':'remove_outliers','column':'x'}])),4)
        self.assertLess(transform(df,[{'kind':'cap_outliers','column':'x'}])['x'].max(),100)
    def test_formula_export(self):
        df=pd.DataFrame({'=header':['=1+1',' +cmd','@SUM(A1)',-4]})
        buf,_=export_dataframe_bytes(df,'xlsx');ws=load_workbook(buf).active
        self.assertNotEqual(ws['A1'].data_type,'f');self.assertNotEqual(ws['A2'].data_type,'f');self.assertEqual(ws['A5'].value,-4)
        csv,_=export_dataframe_bytes(df,'csv');self.assertIn("'=1+1",csv.getvalue().decode())
    def test_csv_json_formats(self):
        df,_,_=parse_uploaded_file('name,value\nJosé,2\n'.encode(),'data.csv');self.assertEqual(df.iloc[0,0],'José')
        df,_,_=parse_uploaded_file(b'[{"a":{"b":2},"arr":[1,2]}]','data.json');self.assertIn('a.b',df)
        with self.assertRaises(ValueError):parse_uploaded_file(b'{}','data.json')
        with self.assertRaises(ValueError):parse_uploaded_file(b'x','data.exe')
    def test_store_expiry_capacity(self):
        s=DatasetStore(ttl=.001);key=s.save(pd.DataFrame({'a':[1]}),'a.csv');time.sleep(.002);self.assertIsNone(s.get(key))
        s=DatasetStore(max_bytes=1)
        with self.assertRaises(ValueError):s.save(pd.DataFrame({'a':[1]}),'a.csv')
    def test_parser_limits(self):
        with patch('server.storage.MAX_ROWS',2):
            with self.assertRaises(ValueError):parse_uploaded_file(b'a\n1\n2\n3','a.csv')
        with self.assertRaises(ValueError):parse_uploaded_file(b'a\x00b','a.csv')

class ApiTests(unittest.TestCase):
    def setUp(self):
        store._datasets.clear();rate_windows.clear();self.client=TestClient(app,raise_server_exceptions=False)
    def upload(self,content=b'name,value\na,1\na,1\nb,3\n',name='test.csv',mime='text/csv'):
        response=self.client.post('/api/upload',files={'file':(name,content,mime)})
        self.assertEqual(response.status_code,200,response.text);return response.json()
    def test_upload_preview_search_and_download(self):
        d=self.upload();key=d['dataset_id'];self.assertGreater(len(key),30)
        r=self.client.get(f'/api/preview/{key}',params={'search':'b'}).json();self.assertEqual(r['total_records'],1)
        for fmt in ['csv','xlsx','json']:self.assertEqual(self.client.get(f'/api/export/{key}?format={fmt}').status_code,200)
        self.assertEqual(self.client.get(f'/api/export/{key}?format=exe').status_code,422)
    def test_ownership_and_delete(self):
        key=self.upload()['dataset_id']
        other=TestClient(app);self.assertEqual(other.get(f'/api/analysis/{key}').status_code,404)
        self.assertEqual(self.client.delete(f'/api/datasets/{key}').status_code,200)
        self.assertEqual(self.client.get(f'/api/analysis/{key}').status_code,404)
    def test_preview_apply_undo_reset(self):
        d=self.upload();key=d['dataset_id'];req={'dataset_id':key,'revision':0,'operations':[{'kind':'duplicates'}]}
        preview=self.client.post('/api/clean',json=req).json();self.assertEqual(preview['after']['rows'],2)
        self.assertEqual(self.client.get(f'/api/preview/{key}').json()['total_records'],3)
        applied=self.client.post('/api/clean',json={**req,'action':'apply'}).json();self.assertEqual(applied['analysis']['summary']['total_rows'],2)
        self.assertEqual(self.client.post('/api/clean',json={**req,'action':'apply'}).status_code,409)
        undone=self.client.post('/api/clean',json={**req,'revision':1,'action':'undo'}).json();self.assertEqual(undone['analysis']['summary']['total_rows'],3)
        reset=self.client.post('/api/clean',json={**req,'revision':2,'action':'reset'});self.assertEqual(reset.status_code,200)
    def test_multisheet(self):
        buf=io.BytesIO()
        with pd.ExcelWriter(buf,engine='openpyxl') as w:
            pd.DataFrame({'x':[1]}).to_excel(w,sheet_name='First',index=False)
            pd.DataFrame({'y':[2,3]}).to_excel(w,sheet_name='Second',index=False)
        d=self.upload(buf.getvalue(),'multi.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        r=self.client.post('/api/switch-sheet',json={'dataset_id':d['dataset_id'],'sheet_name':'Second'})
        self.assertEqual(r.status_code,200,r.text);self.assertEqual(r.json()['analysis']['summary']['total_rows'],2)
        self.assertEqual(r.json()['active_sheet'],'Second')
        self.assertEqual(self.client.post('/api/switch-sheet',json={'dataset_id':d['dataset_id'],'sheet_name':'Missing'}).status_code,400)
    def test_invalid_uploads(self):
        for name,content,mime in [('a.csv',b'','text/csv'),('a.csv',b'a,b\n','text/csv'),('a.exe',b'123','text/plain'),('a.csv',b'a\n1','image/png'),('a.xlsx',b'notzip','application/zip')]:
            r=self.client.post('/api/upload',files={'file':(name,content,mime)});self.assertIn(r.status_code,[400,415],r.text)
            self.assertNotIn('Traceback',r.text)
    def test_large_upload(self):
        r=self.client.post('/api/upload',files={'file':('large.csv',b'a'*(10*1024*1024+100000),'text/csv')});self.assertEqual(r.status_code,413)
    def test_safe_errors(self):
        d=self.upload();r=self.client.post('/api/clean',json={'dataset_id':d['dataset_id'],'revision':0,'operations':[{'kind':'convert','column':'name','value':'number'}]})
        self.assertEqual(r.status_code,400);self.assertNotIn('Unable to parse',r.text)
    def test_rate_limit(self):
        rate_windows['testclient']=(time.monotonic(),120)
        self.assertEqual(self.client.get('/api/health').status_code,429)
    def test_date_and_precision(self):
        d=self.upload(b'date,value\n2025-01-01,1.123456789\n2025-02-01,2.5')
        self.assertEqual(d['analysis']['column_profiles'][0]['type'],'datetime')
        rows=self.client.get('/api/preview/'+d['dataset_id']).json()['records'];self.assertEqual(rows[0]['value'],1.123456789)
    def test_session_listing_and_cross_origin(self):
        d=self.upload()
        self.assertEqual(self.client.get('/api/datasets').json()['datasets'][0]['dataset_id'],d['dataset_id'])
        other=TestClient(app);self.assertEqual(other.get('/api/datasets').json()['datasets'],[])
        r=self.client.post('/api/sample-data',headers={'origin':'https://untrusted.example'})
        self.assertEqual(r.status_code,403)
    def test_invalid_operation_schema(self):
        d=self.upload()
        r=self.client.post('/api/clean',json={'dataset_id':d['dataset_id'],'revision':0,'operations':[{'kind':'fill','method':['malformed']}]})
        self.assertEqual(r.status_code,422);self.assertNotIn('malformed',r.text)
    def test_csv_headers_and_shape(self):
        for body in [b'a,a\n1,2',b'a,b\n1,2,3',b',b\n1,2']:
            self.assertEqual(self.client.post('/api/upload',files={'file':('test.csv',body,'text/csv')}).status_code,400)
    def test_zip_expansion(self):
        payload=io.BytesIO()
        with zipfile.ZipFile(payload,'w',zipfile.ZIP_DEFLATED) as z:
            z.writestr('[Content_Types].xml',b'x'*(51*1024*1024))
        self.assertEqual(self.client.post('/api/upload',files={'file':('bomb.xlsx',payload.getvalue(),'application/zip')}).status_code,400)

    def test_static_routes(self):
        self.assertEqual(self.client.get('/',follow_redirects=False).status_code,200)
        self.assertEqual(self.client.get('/csv-analyzer',follow_redirects=False).status_code,308)
        r=self.client.get('/tools/csv-analyzer');self.assertEqual(r.status_code,200)
        self.assertIn('canonical',r.text);self.assertIn('FAQPage',r.text)
        self.assertEqual(self.client.get('/missing-tool').status_code,404)

if __name__=='__main__':unittest.main()
